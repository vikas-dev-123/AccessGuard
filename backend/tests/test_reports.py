from io import BytesIO

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from tests.test_api_upload import upload_sample

SHEETS = ["Summary", "Terminated Active Access", "Late Revocation", "Orphan Accounts",
          "Generic Shared Accounts", "Privileged Access", "SoD Conflicts", "Dormant Accounts", "All Findings"]


@pytest.fixture
def review_id(client, auditor_headers):
    upload_sample(client, auditor_headers)
    response = client.post("/reviews/run", json={"as_of_date": "2026-09-27"}, headers=auditor_headers)
    return response.json()["id"]


def download(client, headers, review_id, kind):
    response = client.get(f"/reviews/{review_id}/report/{kind}", headers=headers)
    assert response.status_code == 200, response.text
    return response


def test_excel_report_structure(client, viewer_headers, review_id):
    response = download(client, viewer_headers, review_id, "excel")
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert 'filename="AccessGuard_UAR_Apex_Bank_2026-09-27_review' in response.headers["content-disposition"]

    wb = load_workbook(BytesIO(response.content))
    assert wb.sheetnames == SHEETS

    summary = [[c for c in row if c is not None] for row in wb["Summary"].iter_rows(values_only=True)]
    assert ["Total findings", 159] in summary
    assert ["High", 86, 86 / 159] in summary
    assert ["Terminated users with active access", 22, 0, 0, 0, 22, "Terminated Active Access"] in summary

    all_findings = wb["All Findings"]
    assert all_findings.max_row == 160
    assert [c.value for c in all_findings[1]][:3] == ["Finding ID", "Risk", "Check"]

    sod = wb["SoD Conflicts"]
    labels = [sod.cell(row=r, column=1).value for r in range(4, 8)]
    assert labels == ["Condition", "Criteria", "Risk / Impact", "Recommendation"]
    assert sod.cell(row=4, column=2).value.startswith("As of 27 Sep 2026, 19 cases were identified")
    assert len(sod.tables["tbl_sod_conflicts"].ref.split(":")) == 2


def test_pdf_report_content(client, viewer_headers, review_id):
    response = download(client, viewer_headers, review_id, "pdf")
    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF")

    reader = PdfReader(BytesIO(response.content))
    pages = [page.extract_text() for page in reader.pages]
    text = "\n".join(pages)

    assert "User Access Review — Apex Bank" in pages[0]
    assert "Core Banking (100 accounts)" in pages[0]
    assert "1. Executive summary" in text
    assert "159 exceptions were identified: 86 High, 44 Medium, 18 Low and 11 Informational" in text
    for label in ("Condition", "Criteria", "Risk / Impact", "Recommendation"):
        assert label in text
    assert "Appendix A - Full exception list" in text
    assert "F-0159" in text
    assert f"Page {len(pages)} of {len(pages)}" in pages[-1]
    assert reader.pages[-1].mediabox.width > reader.pages[-1].mediabox.height  # landscape appendix


def test_reports_for_review_without_findings(client, auditor_headers):
    import io

    import pandas as pd

    from app.data_loading import ACCOUNT_REQUIRED_COLUMNS, HR_REQUIRED_COLUMNS
    from tests.helpers import account, employee

    def csv(rows, cols):
        return io.BytesIO(pd.DataFrame(rows, columns=cols).to_csv(index=False).encode())

    files = {
        "hr_file": ("hr.csv", csv([employee("E1")], HR_REQUIRED_COLUMNS), "text/csv"),
        "loan_system_file": ("loan.csv", csv([account("L1", "E1", "Credit Analyst")], ACCOUNT_REQUIRED_COLUMNS), "text/csv"),
    }
    assert client.post("/upload", files=files, headers=auditor_headers).status_code == 201
    rid = client.post("/reviews/run", json={"as_of_date": "2026-09-27"}, headers=auditor_headers).json()["id"]

    wb = load_workbook(BytesIO(download(client, auditor_headers, rid, "excel").content))
    assert wb["All Findings"]["A1"].value == "No exceptions noted."
    pdf_text = "\n".join(p.extract_text() for p in PdfReader(BytesIO(download(client, auditor_headers, rid, "pdf").content)).pages)
    assert "No exceptions were identified" in pdf_text
    assert "Appendix A" not in pdf_text


def test_report_requires_auth_and_existing_review(client, viewer_headers):
    assert client.get("/reviews/1/report/pdf").status_code == 401
    assert client.get("/reviews/999/report/excel", headers=viewer_headers).status_code == 404
